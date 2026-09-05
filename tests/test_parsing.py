import io

import pytest
from docx import Document

from ai_rfp_generator.parsing import UnsupportedFileTypeError, extract_text


def test_txt_extraction_decodes_utf8():
    assert extract_text("req.txt", "hello world".encode("utf-8")) == "hello world"


def test_docx_extraction_reads_paragraph_text():
    doc = Document()
    doc.add_paragraph("First requirement")
    doc.add_paragraph("Second requirement")
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_text("req.docx", buffer.getvalue())
    assert "First requirement" in text
    assert "Second requirement" in text


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("req.xlsx", b"whatever")
