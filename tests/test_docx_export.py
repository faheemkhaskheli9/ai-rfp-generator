"""Tests for finalized DOCX export (issue #13)."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from ai_rfp_generator.db import (
    DraftSection,
    Outline,
    OutlineSection,
    Requirement,
    ValidationFinding,
    now_utc,
)


def _seed_export_case(app_module):
    with app_module._SessionFactory() as session:
        requirement = Requirement(
            source_filename="rfp.txt",
            submitted_at=now_utc(),
            status="parsed",
            content="Describe solution and security.",
        )
        session.add(requirement)
        session.flush()

        outline = Outline(
            requirement_id=requirement.id,
            model="test",
            generated_at=now_utc(),
            status="approved",
            reviewed_at=now_utc(),
        )
        session.add(outline)
        session.flush()

        security = OutlineSection(
            outline_id=outline.id,
            position=1,
            title="Security",
            description="Describe security controls.",
        )
        solution = OutlineSection(
            outline_id=outline.id,
            position=0,
            title="Proposed Solution",
            description="Describe the proposed solution.",
        )
        session.add_all([security, solution])
        session.flush()

        solution_draft = DraftSection(
            requirement_id=requirement.id,
            outline_section_id=solution.id,
            strategy="detailed",
            model="test",
            content="Our proposed solution addresses the requested capabilities.",
            fact_ids_json="[]",
            generated_at=now_utc(),
            status="approved",
            reviewed_at=now_utc(),
        )
        security_draft = DraftSection(
            requirement_id=requirement.id,
            outline_section_id=security.id,
            strategy="concise",
            model="test",
            content="Security controls are documented in the proposal.",
            fact_ids_json="[]",
            generated_at=now_utc(),
            status="approved",
            reviewed_at=now_utc(),
        )
        session.add_all([solution_draft, security_draft])
        session.flush()

        finding = ValidationFinding(
            draft_section_id=security_draft.id,
            finding_type="uncited_claim",
            offending_text="Security controls are documented in the proposal.",
            detail="Claim has no source citation.",
            created_at=now_utc(),
        )
        session.add(finding)
        session.commit()
        return outline.id


def test_docx_export_preserves_outline_order_and_headings(client):
    import ai_rfp_generator.app as app_module

    outline_id = _seed_export_case(app_module)

    response = client.post(f"/outlines/{outline_id}/export?format=docx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.content

    document = Document(BytesIO(response.content))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 1"
    ]
    assert headings == ["Table of Contents", "Proposed Solution", "Security"]

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "RFP Response" in text
    assert "Table of Contents" in text
    assert "1. Proposed Solution" in text
    assert "2. Security" in text
    assert "VALIDATION WARNING:" in text
    assert "Our proposed solution addresses the requested capabilities." in text


def test_docx_export_requires_approved_draft(client):
    import ai_rfp_generator.app as app_module

    outline_id = _seed_export_case(app_module)

    with app_module._SessionFactory() as session:
        outline = session.get(Outline, outline_id)
        for section in outline.sections:
            for draft in section.drafts:
                draft.status = "draft"
        session.commit()

    response = client.post(f"/outlines/{outline_id}/export?format=docx")

    assert response.status_code == 422
    assert "no approved section drafts" in response.json()["detail"]


def test_export_rejects_unapproved_outline(client):
    import ai_rfp_generator.app as app_module

    outline_id = _seed_export_case(app_module)
    with app_module._SessionFactory() as session:
        outline = session.get(Outline, outline_id)
        outline.status = "draft"
        session.commit()

    response = client.post(f"/outlines/{outline_id}/export?format=docx")

    assert response.status_code == 422
    assert "outline must be approved" in response.json()["detail"]
