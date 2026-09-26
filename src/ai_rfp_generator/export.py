"""DOCX export for finalized RFP responses."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from ai_rfp_generator.db import DraftSection, Outline

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ExportError(RuntimeError):
    """A proposal cannot be exported in its current state."""


def approved_draft_for_section(section) -> DraftSection | None:
    approved = [draft for draft in section.drafts if draft.status == "approved"]
    if len(approved) > 1:
        raise ExportError(
            f"outline section {section.id} has multiple approved draft versions"
        )
    return approved[0] if approved else None


def build_docx(outline: Outline) -> bytes:
    """Build a Word document from approved draft versions in outline order.

    Sections without an approved draft are omitted. Approved drafts with
    persisted validation findings are included but clearly marked so reviewers
    cannot mistake them for validation-clean content.
    """
    if outline.status != "approved":
        raise ExportError("outline must be approved before export")

    document = Document()
    exported = 0

    for section in sorted(outline.sections, key=lambda item: item.position):
        draft = approved_draft_for_section(section)
        if draft is None:
            continue

        document.add_heading(section.title, level=1)

        if draft.validation_findings:
            warning = document.add_paragraph()
            warning.add_run("VALIDATION WARNING: ").bold = True
            warning.add_run(
                f"{len(draft.validation_findings)} unresolved validation finding(s)."
            )

        for paragraph_text in [part.strip() for part in draft.content.split("\n") if part.strip()]:
            document.add_paragraph(paragraph_text)

        exported += 1

    if exported == 0:
        raise ExportError("no approved section drafts are available for export")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
